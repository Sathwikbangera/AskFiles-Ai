"""Turn an uploaded file's raw bytes into chunked LangChain Documents with
source + page metadata, so answers can cite exactly where they came from."""

import io

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

_splitter = RecursiveCharacterTextSplitter(
    chunk_size=1000,
    chunk_overlap=150,
)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


def _extract_pages(filename: str, content: bytes) -> list[tuple[str, int]]:
    """Returns a list of (page_text, page_number) tuples, page_number is 1-indexed."""
    ext = filename.lower().rsplit(".", 1)[-1]

    if ext == "pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        return [
            (page.extract_text() or "", i + 1) for i, page in enumerate(reader.pages)
        ]

    if ext == "docx":
        from docx import Document as DocxDocument

        doc = DocxDocument(io.BytesIO(content))
        text = "\n".join(p.text for p in doc.paragraphs)
        return [(text, 1)]

    if ext in ("txt", "md"):
        return [(content.decode("utf-8", errors="ignore"), 1)]

    raise ValueError(f"Unsupported file type: .{ext}")


def chunk_file(filename: str, content: bytes) -> list[Document]:
    pages = _extract_pages(filename, content)
    documents = []

    for page_text, page_number in pages:
        if not page_text.strip():
            continue
        for chunk in _splitter.split_text(page_text):
            documents.append(
                Document(
                    page_content=chunk,
                    metadata={"source": filename, "page": page_number},
                )
            )

    return documents
